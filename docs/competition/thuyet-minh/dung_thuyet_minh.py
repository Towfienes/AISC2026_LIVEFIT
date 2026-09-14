"""Dựng file thuyết minh dự thi (.docx) từ nội dung trong ``noi-dung.json``.

Vì sao có tệp này: bản nộp cho ban tổ chức phải DỰNG LẠI ĐƯỢC, y như mọi con
số khác trong dự án. Nội dung nằm ở ``noi-dung.json`` (mảng section, mỗi phần
tử có heading/level/body), định dạng nằm ở đây, còn trang bìa và Phần I lấy
nguyên từ file mẫu của ban tổ chức. Sửa chữ thì sửa JSON rồi chạy lại — sửa
thẳng vào .docx là lần dựng sau mất hết.

Giữ đúng quy định của mẫu: Times New Roman 13, giãn dòng 1,5, lề mặc định của
Word. Khoảng cách giữa các đoạn thì siết lại (mẫu không ràng buộc khoản này)
để bài vừa 29 trang trong giới hạn 15–30.

Quy ước trong ``body``::

    đoạn thường      một dòng văn xuôi
    gạch đầu dòng    "- nội dung"
    ô đánh dấu       "- [x] nội dung"  /  "- [ ] nội dung"
    bảng             "| ô | ô |", hàng đầu là tiêu đề
    sơ đồ ASCII      đặt giữa hai dòng ```
    trong câu        **đậm**, *nghiêng*, `mã`

Cần ``python-docx``, KHÔNG nằm trong phụ thuộc của dự án vì mỗi mùa thi chỉ
dùng một lần. Dựng môi trường riêng rồi chạy::

    python -m venv .venv-docx
    .venv-docx/Scripts/pip install python-docx
    .venv-docx/Scripts/python docs/competition/thuyet-minh/dung_thuyet_minh.py

Đếm số trang (cần Word trên máy)::

    $w = New-Object -ComObject Word.Application
    $d = $w.Documents.Open($duongdan, $false, $true)
    $d.Repaginate(); $d.ComputeStatistics(2)
"""

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

DAY = Path(__file__).resolve().parent
NOI_DUNG = DAY / "noi-dung.json"

# File mẫu của ban tổ chức và nơi ghi bản nộp nằm ở thư mục CHA của kho mã
# (docs/competition/thuyet-minh → docs/competition → docs → kho → thư mục cha).
# Chúng không vào git: mẫu là tài sản của ban tổ chức, còn bản nộp là thứ dựng
# lại được từ tệp này cộng noi-dung.json.
NGOAI = DAY.parents[3]
MAU = NGOAI / "AISC26_Mau_Thuyet_Minh_Du_An.docx"

# Đường ra mặc định; truyền một đường dẫn khác làm tham số dòng lệnh để ghi
# chỗ khác. Cần thật: Word khoá tệp đang mở, nên lúc đang xem bản cũ mà muốn
# dựng bản mới thì ghi ra tệp tạm rồi đổi chỗ sau, thay vì phải đóng Word.
RA = Path(sys.argv[1]) if len(sys.argv) > 1 else NGOAI / "AISC26_Thuyet_Minh_LiveLift.docx"

FONT = "Times New Roman"
CO = Pt(13)
CO_BANG = Pt(11)

THANHVIEN = [
    ("Ngô Bình Minh", "524H0169", "Công nghệ thông tin", "Đại học Tôn Đức Thắng",
     "ngobinhminh2322006@gmail.com", "23/02/2006"),
    ("Lê Xuân Khánh", "524H0055", "Công nghệ thông tin", "Đại học Tôn Đức Thắng",
     "khanhle14062006@gmail.com", "14/06/2006"),
    ("Ngô Lâm Tiến", "524H0195", "Công nghệ thông tin", "Đại học Tôn Đức Thắng",
     "ngolamtien1706@gmail.com", "17/06/2006"),
]
TEN_DOI = "LiveLift"
TRUONG_NHOM = "Ngô Bình Minh"
SDT = "0905484286"

# KHÔNG có cơ chế vá số ở đây, và đó là chủ ý.
#
# Bản dựng ngày 14/09/2026 từng có một danh sách `SUA_SO` vá vài con số lúc
# ghi ra .docx (990→993, 38→41, 8→9 migration). Kết quả: file nộp đúng còn
# `noi-dung.json` sai, tức nguồn và bản dựng nói hai chuyện khác nhau, và lần
# dựng sau sẽ lặng lẽ quay về số cũ. Vòng chấm hồ sơ cùng ngày bắt đúng lỗi đó.
#
# Luật từ nay: MỌI con số sửa thẳng trong `noi-dung.json`. Tệp này chỉ làm
# định dạng, không được đụng tới chữ nghĩa.


def set_font(run, *, bold=False, size=CO, italic=False, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    run.font.size = Pt(9) if mono else size
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    return run


def para(doc, text="", *, bold=False, size=CO, italic=False, align=None,
         space_before=0, space_after=6, indent=None, mono=False, spacing=1.5,
         giu_voi_doan_sau=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0 if mono else spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Pt(indent)
    if mono:
        # Sơ đồ ASCII: bỏ thụt dòng đầu của docDefaults, nếu không dòng đầu
        # khối bị đẩy lệch 1 cm so với các dòng sau và khung vẽ vỡ.
        pf.first_line_indent = Pt(0)
    if giu_voi_doan_sau:
        # Tiêu đề không được nằm trơ một mình ở dòng cuối trang.
        pf.keep_with_next = True
    if text:
        set_font(p.add_run(text), bold=bold, size=size, italic=italic, mono=mono)
    return p


def emit_runs(p, text, *, size=CO, bold_default=False):
    """Dịch **đậm**, *nghiêng* và `mã` thành run Word.

    Phải bắt cả dấu sao ĐƠN: tên tạp chí trong mục tài liệu tham khảo viết kiểu
    ``*Management Science*``, và nếu không xử thì dấu sao in ra thành ký tự
    thật giữa câu — lỗi này đã lọt tới bản PDF một lần rồi.
    """
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            set_font(p.add_run(chunk[2:-2]), bold=True, size=size)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            set_font(p.add_run(chunk[1:-1]), italic=True, size=size)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            set_font(p.add_run(chunk[1:-1]), mono=True)
        else:
            set_font(p.add_run(chunk), bold=bold_default, size=size)


def rich_para(doc, text, *, indent=None, space_after=3, hanging=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    # Văn xuôi canh đều cho đẹp; gạch đầu dòng canh trái, vì canh đều trên
    # dòng có thụt treo làm chữ giãn ra thành khoảng trắng loang lổ.
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if hanging else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent is not None:
        pf.left_indent = Pt(indent)
        if hanging:
            pf.first_line_indent = Pt(-12)
    emit_runs(p, text)
    return p


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cp = cells[j].paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_after = Pt(1)
            cp.paragraph_format.space_before = Pt(1)
            # Mẫu để mặc định canh giữa; ô nhiều chữ canh giữa rất khó đọc.
            cp.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i == 0
                            else WD_ALIGN_PARAGRAPH.LEFT)
            # PHẢI đặt tường minh về 0. File mẫu đặt `<w:ind w:firstLine="567"/>`
            # ở docDefaults, tức thụt dòng đầu 1 cm cho MỌI đoạn — kể cả đoạn
            # nằm trong ô bảng. Với văn xuôi thì đó là thụt đầu dòng bình thường
            # và đẹp, nhưng trong ô bảng nó đẩy dòng đầu vào trong còn dòng
            # xuống hàng lại sát mép trái, trông như lỗi in. Vòng chấm hồ sơ
            # 14/09/2026 bắt lỗi này ở cả 22 bảng.
            cp.paragraph_format.first_line_indent = Pt(0)
            emit_runs(cp, txt, size=CO_BANG, bold_default=(i == 0))

        # Hàng tiêu đề lặp lại ở mỗi trang, và không hàng nào bị xẻ đôi giữa
        # hai trang: bảng dài mà mất tiêu đề thì người chấm phải lật ngược
        # lại mới biết cột nào là cột nào.
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        khong_xe = OxmlElement("w:cantSplit")
        trpr.append(khong_xe)
    return t


def render_body(doc, body):
    lines = body.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # sơ đồ ASCII: giữ khoảng trắng, phông đều nét
        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                bl = lines[i].rstrip()
                para(doc, bl if bl else " ", mono=True, space_after=0, indent=10)
                i += 1
            i += 1
            para(doc, "", space_after=2)
            continue

        # bảng markdown
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cur = lines[i].strip()
                if not re.fullmatch(r"[\s|:\-]+", cur):
                    rows.append([c.strip() for c in cur.strip("|").split("|")])
                i += 1
            if rows:
                add_table(doc, rows)
                para(doc, "", space_after=2)
            continue

        if not stripped:
            i += 1
            continue

        # ô đánh dấu — phải xét TRƯỚC bullet vì viết dạng "- [x] ..."
        m = re.match(r"^\s*(?:[-•]\s*)?\[([ xX])\]\s+(.*)$", line)
        if m:
            mark = "☒" if m.group(1).lower() == "x" else "☐"
            rich_para(doc, f"{mark}  {m.group(2)}", indent=18, space_after=2, hanging=True)
            i += 1
            continue

        # gạch đầu dòng
        m = re.match(r"^(\s*)[-•]\s+(.*)$", line)
        if m:
            depth = min(len(m.group(1)) // 2, 2)
            rich_para(doc, "• " + m.group(2), indent=18 + depth * 16,
                      space_after=2, hanging=True)
            i += 1
            continue

        # danh sách đánh số
        if re.match(r"^\s*\d+[.)]\s", line):
            rich_para(doc, stripped, indent=18, space_after=2, hanging=True)
            i += 1
            continue

        rich_para(doc, stripped)
        i += 1


def dien_phan_i(doc):
    for p in doc.paragraphs:
        t = p.text.strip()
        moi = None
        if t == "Tên đội thi:" or t.startswith("Tên đội thi:"):
            moi = f"Tên đội thi: {TEN_DOI}"
        elif t == "Họ và tên:":
            moi = f"Họ và tên: {TRUONG_NHOM}"
        elif t.startswith("Số điện thoại:"):
            moi = f"Số điện thoại: {SDT}"
        if moi is None:
            continue
        keep_bold = any(r.font.bold for r in p.runs)
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        set_font(p.add_run(moi), bold=bool(keep_bold))

    tbl = next((t for t in doc.tables
                if t.rows and "Thành viên" in "".join(c.text for c in t.rows[0].cells)), None)
    if tbl is None:
        sys.exit("Không tìm thấy bảng thành viên trong mẫu")
    for ri in range(1, min(7, len(tbl.rows))):
        row = tbl.rows[ri]
        for mi in range(5):
            ci = mi + 1
            if ci >= len(row.cells):
                break
            val = THANHVIEN[mi][ri - 1] if mi < len(THANHVIEN) else ""
            cell = row.cells[ci]
            cell.text = ""
            cp = cell.paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            set_font(cp.add_run(val), size=CO_BANG)


def main():
    secs = json.loads(NOI_DUNG.read_text(encoding="utf-8"))
    doc = Document(str(MAU))
    dien_phan_i(doc)

    # Xoá nội dung Phần II cũ của mẫu
    body = doc.element.body
    start = None
    for idx, child in enumerate(body):
        if child.tag == qn("w:p"):
            txt = "".join(n.text or "" for n in child.iter(qn("w:t")))
            if "PHẦN II" in txt.upper():
                start = idx
                break
    if start is None:
        sys.exit("Không tìm thấy mốc PHẦN II trong mẫu")
    for child in list(body)[start + 1:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    # Giữ nguyên w:sectPr: python-docx cần nó để tính bề rộng bảng, và tự chèn
    # mọi đoạn/bảng mới vào TRƯỚC thẻ này nên thứ tự vẫn đúng.
    for s in secs:
        if s.get("level", 1) == 1:
            para(doc, s["heading"].upper(), bold=True, size=Pt(14),
                 space_before=10, space_after=4, giu_voi_doan_sau=True)
        else:
            para(doc, s["heading"], bold=True, size=Pt(13),
                 space_before=8, space_after=4, giu_voi_doan_sau=True)
        render_body(doc, s["body"])

    doc.save(str(RA))
    print("Đã ghi:", RA)
    print("Số đoạn:", len(doc.paragraphs), "| số bảng:", len(doc.tables))


if __name__ == "__main__":
    main()
