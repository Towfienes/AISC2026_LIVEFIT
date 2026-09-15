"""Dựng hồ sơ dự án Bảng C (.docx) cho Cuộc thi Sáng tạo trẻ Quốc gia về AI 2026.

Vì sao có tệp này: bản nộp phải DỰNG LẠI ĐƯỢC. Nội dung nằm ở ``noi-dung.md``,
định dạng nằm ở đây, còn khung MẪU 3 lấy NGUYÊN XI từ file mẫu của ban tổ chức
— kể cả bảng thông tin thí sinh và khối ký tên. Sửa chữ thì sửa Markdown rồi
chạy lại; sửa thẳng .docx là lần dựng sau mất hết.

Bộ dựng AISC dùng JSON; tệp này dùng Markdown vì hồ sơ phải qua tay ba người
trong hai tuần, và sửa văn xuôi tiếng Việt trong JSON một dòng thì ai cũng
làm hỏng dấu ngoặc. ``# `` là mục cấp 1, ``## `` cấp 2, ``### `` cấp 3.

Khác với bộ dựng của AISC ở ``docs/competition/thuyet-minh/``: mẫu của cuộc thi
này gộp cả ba bảng A, B, C vào MỘT file. Tệp này cắt lấy đúng đoạn MẪU 3 (Bảng
C) rồi mới điền, nên nếu ban tổ chức phát hành lại mẫu thì chỉ cần thay file
mẫu, không phải sửa mã.

Giới hạn cứng: **20 trang**. Mẫu để giãn dòng 1,15 nên bộ dựng giữ đúng 1,15
thay vì 1,5 — vừa đúng mẫu vừa đủ chỗ cho 13 mục. Chạy ``--dem-trang`` để Word
đếm hộ trước khi nộp.

Cần ``python-docx``, cố ý KHÔNG nằm trong phụ thuộc dự án::

    python -m venv .venv-docx
    .venv-docx/Scripts/pip install python-docx
    .venv-docx/Scripts/python docs/competition/sang-tao-tre-2026/dung_ho_so.py

Quy ước trong ``body`` (giống bộ dựng AISC, để chép qua lại được)::

    đoạn thường      một dòng văn xuôi
    gạch đầu dòng    "- nội dung"
    bảng             "| ô | ô |", hàng đầu là tiêu đề
    sơ đồ ASCII      đặt giữa hai dòng ```
    ảnh              "![chú thích](đường/dẫn/anh.png)"
    trong câu        **đậm**, *nghiêng*, `mã`
"""

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DAY = Path(__file__).resolve().parent
NOI_DUNG = DAY / "noi-dung.md"

# Mẫu của ban tổ chức nằm ngoài kho mã (tài sản của BTC, không đưa vào git).
MAU = Path(
    r"D:\AISC2026\sang tạo trẻ quốc gia\drive-download-20260914T062853Z-1-001"
    r"\AI2026_Mẫu hồ sơ.docx"
)
RA_MAC_DINH = Path(r"D:\AISC2026\AI2026_Ho_So_Du_An_LiveLift_BangC.docx")

FONT = "Times New Roman"
CO = Pt(13)
CO_BANG = Pt(10.5)
GIAN_DONG = 1.15  # đúng như mẫu MẪU 3 đặt cho phần nội dung

# ---------------------------------------------------------------- thông tin đội
# Đọc từ docs/competition/thong-tin-doi.local.json (đã gitignore) — ngày sinh,
# MSSV, số điện thoại không được nằm trong kho mã công khai. Xem thong_tin_doi.py.
# Ô nào còn "⬜" là đội PHẢI tự điền trước khi nộp.
sys.path.insert(0, str(DAY.parent))
import thong_tin_doi  # noqa: E402

_DOI = thong_tin_doi.doc()
THANH_VIEN = [
    {
        "ho_ten": tv["ho_ten"],
        "ngay_sinh": tv["ngay_sinh"],
        "lop": f"Ngành {tv['nganh']}, Khoa {tv['khoa']}, Trường {tv['truong']} (MSSV {tv['mssv']})",
        "noi_o": tv["noi_o"],
        "dien_thoai": tv["dien_thoai"],
        "email": tv["email"],
    }
    for tv in _DOI["thanh_vien"]
]
SO_THANH_VIEN = len(THANH_VIEN)


# ------------------------------------------------------------------ tiện ích Word
def set_font(run, *, bold=False, size=CO, italic=False, mono=False):
    name = "Consolas" if mono else FONT
    run.font.name = name
    run.font.size = Pt(8.5) if mono else size
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    return run


def emit_runs(p, text, *, size=CO, bold_default=False):
    """Dịch **đậm**, *nghiêng* và `mã` thành run Word.

    Bắt cả dấu sao ĐƠN: tên tạp chí trong mục tài liệu tham khảo viết kiểu
    ``*Production and Operations Management*``; không xử thì dấu sao in ra
    thành ký tự thật giữa câu.
    """
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            set_font(p.add_run(chunk[2:-2]), bold=True, size=size)
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            set_font(p.add_run(chunk[1:-1]), italic=True, size=size)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            set_font(p.add_run(chunk[1:-1]), mono=True)
        else:
            set_font(p.add_run(chunk), bold=bold_default, size=size)


def para(
    doc,
    text="",
    *,
    bold=False,
    size=CO,
    italic=False,
    align=None,
    space_before=0,
    space_after=4,
    indent=None,
    mono=False,
    giu_voi_doan_sau=False,
):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.0 if mono else GIAN_DONG
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Pt(0)
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Pt(indent)
    if giu_voi_doan_sau:
        pf.keep_with_next = True
    if text:
        set_font(p.add_run(text), bold=bold, size=size, italic=italic, mono=mono)
    return p


def rich_para(doc, text, *, indent=None, space_after=3, hanging=False, italic_all=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = GIAN_DONG
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Pt(0)
    # Văn xuôi canh đều; gạch đầu dòng canh trái, vì canh đều trên dòng có
    # thụt treo làm chữ giãn thành khoảng trắng loang lổ.
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if hanging else WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent is not None:
        pf.left_indent = Pt(indent)
        if hanging:
            pf.first_line_indent = Pt(-11)
    if italic_all:
        set_font(p.add_run(text), italic=True)
    else:
        emit_runs(p, text)
    return p


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(ncol):
            txt = row[j] if j < len(row) else ""
            cp = cells[j].paragraphs[0]
            cp.paragraph_format.line_spacing = 1.0
            cp.paragraph_format.space_after = Pt(1)
            cp.paragraph_format.space_before = Pt(1)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            # Mẫu đặt thụt dòng đầu ở docDefaults; trong ô bảng nó đẩy dòng đầu
            # vào trong còn dòng xuống hàng sát mép trái, trông như lỗi in.
            cp.paragraph_format.first_line_indent = Pt(0)
            emit_runs(cp, txt, size=CO_BANG, bold_default=(i == 0))

        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))  # lặp tiêu đề mỗi trang
        trpr.append(OxmlElement("w:cantSplit"))  # không xẻ đôi hàng
    return t


def render_body(doc, body):
    lines = body.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # sơ đồ ASCII
        if stripped.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                bl = lines[i].rstrip()
                para(doc, bl if bl else " ", mono=True, space_after=0, indent=6)
                i += 1
            i += 1
            para(doc, "", space_after=2)
            continue

        # ảnh minh chứng: ![chú thích](đường dẫn) hoặc ![chú thích|11](đường dẫn)
        # Số sau dấu | là bề rộng tính bằng cm. Mặc định 13 cm: ảnh chụp màn
        # hình tỉ lệ 16:10 ở 15,5 cm cao gần 10 cm, ba ảnh là mất hơn một trang
        # trong khi hồ sơ chỉ có 20 trang.
        m = re.match(r"^!\[(.*?)\]\((.+?)\)$", stripped)
        if m:
            chu_thich, duong_dan = m.group(1), m.group(2)
            rong = 13.0
            if "|" in chu_thich:
                chu_thich, _, so = chu_thich.rpartition("|")
                try:
                    rong = float(so.strip())
                except ValueError:
                    chu_thich = m.group(1)
            anh = (DAY / duong_dan) if not Path(duong_dan).is_absolute() else Path(duong_dan)
            if anh.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True  # ảnh không rời chú thích
                p.add_run().add_picture(str(anh), width=Cm(rong))
                cap = rich_para(doc, chu_thich, space_after=6, italic_all=True)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                # Thiếu ảnh thì phải THẤY được, không im lặng bỏ qua.
                rich_para(doc, f"[THIẾU ẢNH: {duong_dan}] {chu_thich}")
            i += 1
            continue

        # bảng markdown
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cur = lines[i].strip()
                if not re.fullmatch(r"[\s|:\-]+", cur):
                    rows.append([c.strip() for c in cur.strip("|").split("|")])
                i += 1
            if rows:
                add_table(doc, rows)
                para(doc, "", space_after=2)
            continue

        if not stripped:
            i += 1
            continue

        # gạch đầu dòng
        m = re.match(r"^(\s*)[-•]\s+(.*)$", line)
        if m:
            depth = min(len(m.group(1)) // 2, 2)
            rich_para(doc, "• " + m.group(2), indent=14 + depth * 14, space_after=2, hanging=True)
            i += 1
            continue

        # danh sách đánh số
        if re.match(r"^\s*\d+[.)]\s", line):
            rich_para(doc, stripped, indent=14, space_after=2, hanging=True)
            i += 1
            continue

        rich_para(doc, stripped)
        i += 1


# ------------------------------------------------------------- cắt và điền mẫu
def cat_lay_mau_3(doc):
    """Xoá mọi thứ không thuộc MẪU 3, giữ nguyên sectPr ở cuối."""
    body = doc.element.body
    con = list(body)
    dau = None
    for idx, c in enumerate(con):
        if c.tag != qn("w:p"):
            continue
        txt = "".join(n.text or "" for n in c.iter(qn("w:t")))
        if "MẪU HỒ SƠ DỰ ÁN DỰ THI BẢNG C" in txt:
            dau = idx
            break
    if dau is None:
        sys.exit("Không tìm thấy MẪU 3 (Bảng C) trong file mẫu của ban tổ chức")
    for c in con[:dau]:
        body.remove(c)
    return dau


def diem_moc_noi_dung(doc):
    """Vị trí đoạn 'NỘI DUNG HỒ SƠ DỰ ÁN' và bảng ký tên cuối mẫu."""
    body = doc.element.body
    con = list(body)
    moc = None
    for idx, c in enumerate(con):
        if c.tag != qn("w:p"):
            continue
        txt = "".join(n.text or "" for n in c.iter(qn("w:t")))
        if "NỘI DUNG HỒ SƠ DỰ ÁN" in txt:
            moc = idx
            break
    if moc is None:
        sys.exit("Không tìm thấy mốc 'NỘI DUNG HỒ SƠ DỰ ÁN'")
    ky_ten = None
    for c in con[moc:]:
        if c.tag == qn("w:tbl"):
            txt = "".join(n.text or "" for n in c.iter(qn("w:t")))
            if "Đại diện đội thi" in txt:
                ky_ten = c
                break
    return moc, ky_ten


def dien_thong_tin_thi_sinh(doc):
    """Điền bảng thông tin 3 thí sinh của MẪU 3.

    Mẫu là MỘT bảng dọc: hàng đầu là 'Số lượng thí sinh', rồi mỗi thí sinh 6
    hàng nhãn-giá trị. Điền theo NHÃN chứ không theo chỉ số hàng, vì ban tổ
    chức có thể phát hành lại mẫu với số hàng khác.
    """
    tbl = next(
        (t for t in doc.tables if "Số lượng thí sinh" in "".join(c.text for c in t.rows[0].cells)),
        None,
    )
    if tbl is None:
        sys.exit("Không tìm thấy bảng thông tin thí sinh trong MẪU 3")

    khoa = {
        "họ và tên": "ho_ten",
        "ngày/tháng/năm sinh": "ngay_sinh",
        "lớp hành chính, ngành, khoa, trường": "lop",
        "xã/phường/đặc khu, tỉnh/thành phố": "noi_o",
        "điện thoại": "dien_thoai",
        "email": "email",
    }
    chi_so = -1  # -1 = chưa tới thí sinh nào
    da_dien = 0
    for row in tbl.rows:
        nhan = row.cells[0].text.strip().lower().replace("\n", " ")
        nhan = re.sub(r"\s+", " ", nhan)

        if nhan.startswith("số lượng thí sinh"):
            # Tích ô "3 người". Các ô ☐ nằm ngay trên cùng hàng.
            for cell in row.cells[1:]:
                if f"{SO_THANH_VIEN} người" in cell.text:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.text = r.text.replace("☐", "☒")
            continue

        if nhan.startswith("thí sinh thứ"):
            chi_so += 1
            continue

        if chi_so < 0 or chi_so >= len(THANH_VIEN):
            continue

        for k, truong in khoa.items():
            if nhan.startswith(k.split(",")[0]) and k.split(",")[0] in nhan:
                cell = row.cells[1] if len(row.cells) > 1 else None
                if cell is None:
                    break
                cell.text = ""
                cp = cell.paragraphs[0]
                cp.paragraph_format.line_spacing = 1.0
                cp.paragraph_format.first_line_indent = Pt(0)
                set_font(cp.add_run(THANH_VIEN[chi_so][truong]), size=CO_BANG)
                da_dien += 1
                break
    if da_dien < len(THANH_VIEN) * 4:
        print(
            f"  CẢNH BÁO: chỉ điền được {da_dien} ô thông tin thí sinh — "
            "kiểm tra lại mẫu của ban tổ chức có đổi nhãn không"
        )
    return da_dien


def doc_muc(duong_dan):
    """Tách Markdown thành [{heading, level, body}].

    Bỏ qua mọi dòng ``<!-- ... -->``: chú thích cho người viết, không ra bản nộp.
    """
    secs = []
    cur = None
    for line in duong_dan.read_text(encoding="utf-8").splitlines():
        if line.lstrip().startswith("<!--"):
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            cur = {"heading": m.group(2).strip(), "level": len(m.group(1)), "body": []}
            secs.append(cur)
            continue
        if cur is not None:
            cur["body"].append(line)
    for s in secs:
        s["body"] = "\n".join(s["body"]).strip("\n")
    if not secs:
        sys.exit(f"{duong_dan.name} không có mục nào (cần dòng bắt đầu bằng '# ')")
    return secs


GIOI_HAN_TRANG = 20


def xuat_pdf_va_dem_trang(docx):
    """Nhờ Word đếm trang rồi xuất PDF. Trả (số trang, số từ, đường dẫn PDF).

    Vì sao phải là Word: giới hạn 20 trang là giới hạn CỨNG, và chỉ Word mới
    ngắt trang giống thứ ban tổ chức sẽ mở ra. LibreOffice và các thư viện
    Python ngắt khác, lệch một hai trang là chuyện thường.

    Lệnh chạy thẳng qua -Command chứ không qua tệp .ps1: Windows PowerShell 5.1
    đọc tệp .ps1 theo bảng mã ANSI nếu tệp không có BOM, nên mọi chú thích
    tiếng Việt trong .ps1 làm hỏng bộ phân tích cú pháp. Đã dính một lần rồi.
    """
    pdf = docx.with_suffix(".pdf")
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx}', $false, $true); $d.Repaginate(); "
        "$t = $d.ComputeStatistics(2); $u = $d.ComputeStatistics(0); "
        f"$d.SaveAs([ref]'{pdf}', [ref]17); $d.Close($false); $w.Quit(); "
        'Write-Output "$t $u"'
    )
    try:
        # noqa S603/S607: lệnh dựng từ hằng số trong tệp này cộng đường dẫn
        # tệp do chính script sinh ra — không có đầu vào từ người dùng.
        out = subprocess.run(  # noqa: S603
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],  # noqa: S607
            capture_output=True,
            text=True,
            timeout=180,
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"  Không đếm trang được ({e}) — mở Word xem thanh trạng thái")
        return None, None, None
    so = out.stdout.split()
    if len(so) < 2 or not so[0].isdigit():
        print("  Không đếm trang được. Word có đang mở chính tệp này không?")
        if out.stderr.strip():
            print("  ", out.stderr.strip().splitlines()[0])
        return None, None, None
    return int(so[0]), int(so[1]), pdf


def main():
    argv = [a for a in sys.argv[1:] if not a.startswith("--")]
    ra = Path(argv[0]) if argv else RA_MAC_DINH

    if not MAU.exists():
        sys.exit(f"Không thấy file mẫu của ban tổ chức: {MAU}")
    secs = doc_muc(NOI_DUNG)

    doc = Document(str(MAU))
    cat_lay_mau_3(doc)
    so_o = dien_thong_tin_thi_sinh(doc)

    moc, ky_ten = diem_moc_noi_dung(doc)
    body = doc.element.body
    if ky_ten is not None:
        body.remove(ky_ten)  # đặt lại ở cuối sau khi đổ nội dung
    for c in list(body)[moc + 1 :]:
        if c.tag != qn("w:sectPr"):
            body.remove(c)

    for s in secs:
        lv = s.get("level", 1)
        if lv == 1:
            para(
                doc,
                s["heading"],
                bold=True,
                size=Pt(13.5),
                space_before=9,
                space_after=3,
                giu_voi_doan_sau=True,
            )
        else:
            para(
                doc,
                s["heading"],
                bold=True,
                size=Pt(13),
                italic=(lv >= 3),
                space_before=6,
                space_after=3,
                giu_voi_doan_sau=True,
            )
        render_body(doc, s["body"])

    if ky_ten is not None:
        para(doc, "", space_before=8, space_after=0)
        sect = body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(ky_ten)
        else:
            body.append(ky_ten)

    ra.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ra))
    print("Đã ghi:", ra)
    print(
        f"  {len(secs)} mục · {len(doc.paragraphs)} đoạn · {len(doc.tables)} bảng "
        f"· {so_o} ô thông tin thí sinh"
    )
    if _DOI["thieu_tep_that"]:
        print(
            "  CẢNH BÁO: chưa có docs/competition/thong-tin-doi.local.json — "
            "đang dựng bằng tệp mẫu, bảng thí sinh toàn ô trống"
        )
    trang, tu, pdf = xuat_pdf_va_dem_trang(ra)
    if trang is not None:
        print(f"  {trang}/{GIOI_HAN_TRANG} trang · {tu} từ · PDF: {pdf}")

    canh_bao = []
    chua_dien = sum(1 for tv in THANH_VIEN for v in tv.values() if "⬜" in v)
    if chua_dien:
        canh_bao.append(
            f"{chua_dien} ô thông tin thí sinh còn dấu ⬜ "
            "(điền docs/competition/thong-tin-doi.local.json)"
        )
    con_trong = sum(1 for s in secs if "⬜" in s["body"])
    if con_trong:
        canh_bao.append(f"{con_trong} mục còn dấu ⬜ trong noi-dung.md")
    if trang is not None and trang > GIOI_HAN_TRANG:
        canh_bao.append(f"VƯỢT GIỚI HẠN {GIOI_HAN_TRANG} TRANG — phải cắt bớt")
    if canh_bao:
        print("  CHƯA NỘP ĐƯỢC:")
        for c in canh_bao:
            print("   -", c)
        return 1
    print("  Sẵn sàng nộp.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
