"""Dựng Bản kê khai công cụ AI / dữ liệu / API / thư viện (DOCX + PDF) từ MỘT nguồn.

Nguồn duy nhất là ``docs/competition/sang-tao-tre-2026/05-BAN-KE-KHAI.md``. Sửa chữ
thì sửa tệp .md rồi chạy lại, KHÔNG sửa thẳng tệp .docx.

    # cần venv riêng có python-docx (đã gitignore): .venv-docx
    .venv-docx/Scripts/python docs/competition/sang-tao-tre-2026/ke_khai/dung_ke_khai.py
    .venv-docx/Scripts/python docs/competition/sang-tao-tre-2026/ke_khai/dung_ke_khai.py \\
        --ra D:/AISC2026/AI2026_Ban_Ke_Khai_LiveLift.docx
    # chỉ .docx, không gọi Word:
    ... dung_ke_khai.py --khong-pdf

Định dạng: A4, lề trái 3 cm, phải/trên/dưới 2 cm; Times New Roman 13 cho thân bài,
11 cho bảng; tiêu đề đậm; bảng có viền, lặp dòng tiêu đề khi sang trang; số trang
"Trang X/Y" ở chân trang; bảng chữ ký cao đủ để ký tay.

Kiểm TRƯỚC khi thay tệp cũ: không còn ký hiệu Markdown (backtick, ``**``, ``](``) và
không còn ô ⬜ trong văn bản. PDF xuất bằng Word (COM) — cùng bộ máy ngắt trang mà
người chấm sẽ mở. Lỗi ở bất kỳ bước nào thì giữ nguyên bản cũ và thoát mã 1.
"""

from __future__ import annotations

import argparse
import contextlib
import os
import re
import subprocess
import sys
from pathlib import Path

DAY = Path(__file__).resolve().parent
sys.path.insert(0, str(DAY))
from doc_md import Khoi, chu_tron, phan_tich, tach_trong_dong  # noqa: E402

NGUON = DAY.parent / "05-BAN-KE-KHAI.md"
RA_MAC_DINH = Path("D:/AISC2026/AI2026_Ban_Ke_Khai_LiveLift.docx")
FONT = "Times New Roman"
FONT_MA = "Consolas"
CO = 13
CO_BANG = 11
NEN_TIEU_DE_BANG = "D9E2F3"
NEN_MA = "F2F2F2"


def _nap_docx():
    try:
        import docx  # noqa: F401
    except ImportError:
        sys.exit(
            "Thiếu python-docx. Chạy bằng venv riêng: "
            ".venv-docx/Scripts/python docs/competition/sang-tao-tre-2026/ke_khai/dung_ke_khai.py"
        )


# ------------------------------------------------------------------ tiện ích Word
def _dat_font(run, ten: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), ten)


def _nen(el_pr, mau: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), mau)
    el_pr.append(shd)


def _run(p, chu, *, co=CO, dam=False, nghieng=False, ma=False, gach=False):
    from docx.shared import Pt

    r = p.add_run(chu)
    ten = FONT_MA if ma else FONT
    r.font.name = ten
    r.font.size = Pt(round(co * 0.85 * 2) / 2) if ma else Pt(co)
    r.font.bold = dam
    r.font.italic = nghieng
    r.font.underline = gach
    _dat_font(r, ten)
    if ma:
        _nen(r._element.get_or_add_rPr(), NEN_MA)
    return r


def _chu_dinh_dang(p, chu: str, *, co=CO, dam=False, nghieng=False) -> None:
    for d in tach_trong_dong(chu, dam=dam, nghieng=nghieng):
        _run(p, d.chu, co=co, dam=d.dam, nghieng=d.nghieng, ma=d.ma)


def _doan(
    doc_or_cell,
    chu="",
    *,
    co=CO,
    dam=False,
    nghieng=False,
    canh=None,
    truoc=0,
    sau=4,
    thut_trai=None,
    thut_dau=None,
    giu_voi_sau=False,
):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(truoc)
    pf.space_after = Pt(sau)
    pf.line_spacing = 1.15
    pf.keep_with_next = giu_voi_sau
    if thut_trai is not None:
        pf.left_indent = Cm(thut_trai)
    if thut_dau is not None:
        pf.first_line_indent = Cm(thut_dau)
    p.alignment = {
        None: WD_ALIGN_PARAGRAPH.JUSTIFY,
        "giua": WD_ALIGN_PARAGRAPH.CENTER,
        "phai": WD_ALIGN_PARAGRAPH.RIGHT,
        "trai": WD_ALIGN_PARAGRAPH.LEFT,
    }[canh]
    if chu:
        _chu_dinh_dang(p, chu, co=co, dam=dam, nghieng=nghieng)
    return p


def _lap_tieu_de_bang(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)
    khong_xe = OxmlElement("w:cantSplit")
    khong_xe.set(qn("w:val"), "true")
    trpr.append(khong_xe)


def _bo_cuc_co_dinh(t) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tblpr = t._tbl.tblPr
    el = tblpr.find(qn("w:tblLayout"))
    if el is None:
        el = OxmlElement("w:tblLayout")
        tblpr.append(el)
    el.set(qn("w:type"), "fixed")


def _khong_xe_hang(row) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def _do_rong_cot(hang: list[list[str]], ncot: int, tong_cm: float) -> list[float]:
    """Mỗi cột đủ rộng cho TỪ dài nhất của nó (không xẻ đôi "19.126" hay "Antigravity"),
    phần còn lại chia theo độ dài trung bình và lớn nhất của ô."""
    cm_moi_ky_tu = 0.21  # Times New Roman 11 pt, xấp xỉ; đường dẫn dài vẫn được Word ngắt
    tu_dai = [1] * ncot
    dai_max = [1] * ncot
    dai_tb = [0.0] * ncot
    for j in range(ncot):
        o = [chu_tron(h[j]) if j < len(h) else "" for h in hang]
        tu_dai[j] = max([len(t) for x in o for t in x.split()] or [1])
        dai_max[j] = max(min(len(x), 80) for x in o) or 1
        dai_tb[j] = sum(min(len(x), 80) for x in o) / max(len(o), 1)
    toi_thieu = [min(max(0.8, min(t, 14) * cm_moi_ky_tu + 0.45), 3.4) for t in tu_dai]
    if sum(toi_thieu) >= tong_cm:
        return [tong_cm * t / sum(toi_thieu) for t in toi_thieu]
    trong_so = [(dai_tb[j] + dai_max[j]) / 2 for j in range(ncot)]
    con_lai = tong_cm - sum(toi_thieu)
    return [toi_thieu[j] + con_lai * trong_so[j] / sum(trong_so) for j in range(ncot)]


def _bang(doc, khoi: Khoi, *, ky_ten=False) -> None:
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    hang = khoi.muc
    if not hang:
        return
    ncot = max(len(h) for h in hang)
    t = doc.add_table(rows=len(hang), cols=ncot)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rong = _do_rong_cot(hang, ncot, 16.0)
    if ky_ten and ncot == 4:  # họ tên · vai trò · tự khai · CHỮ KÝ (đủ rộng để ký tay)
        rong = [3.3, 2.3, 6.4, 4.0]
    # Word tự co giãn cột (autofit) và bỏ qua độ rộng từng ô nếu lưới cột không khớp:
    # tắt autofit, bố cục cố định, ghi độ rộng vào cả lưới cột.
    t.autofit = False
    _bo_cuc_co_dinh(t)
    for j, col in enumerate(t.columns):
        col.width = Cm(rong[j])
    canh_map = {"giua": WD_ALIGN_PARAGRAPH.CENTER, "phai": WD_ALIGN_PARAGRAPH.RIGHT}
    for i, h in enumerate(hang):
        row = t.rows[i]
        if i == 0:
            _lap_tieu_de_bang(row)
        elif ky_ten:
            row.height = Cm(2.8)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            _khong_xe_hang(row)
        if ky_ten:  # giữ cả bảng ký trên một trang
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = i < len(hang) - 1
        for j in range(ncot):
            cell = row.cells[j]
            cell.width = Cm(rong[j])
            o = h[j] if j < len(h) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _nen(cell._tc.get_or_add_tcPr(), NEN_TIEU_DE_BANG)
                _chu_dinh_dang(p, o, co=CO_BANG, dam=True)
            else:
                if j < len(khoi.canh) and khoi.canh[j] in canh_map:
                    p.alignment = canh_map[khoi.canh[j]]
                # <br> trong ô = xuống dòng
                phan = re.split(r"<br\s*/?>", o)
                for k, ph in enumerate(phan):
                    if k:
                        p = cell.add_paragraph()
                        p.paragraph_format.space_after = Pt(1)
                    _chu_dinh_dang(p, ph, co=CO_BANG)
    _doan(doc, "", sau=2)


def _quoc_hieu(doc) -> None:
    _doan(doc, "**CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM**", canh="giua", sau=0)
    p = _doan(doc, "", canh="giua", sau=10)
    _run(p, "Độc lập – Tự do – Hạnh phúc", dam=True, gach=True)


def _so_trang(doc) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "Bản kê khai — LiveLift · Trang ", co=10, nghieng=True)
        for truong in ("PAGE", "NUMPAGES"):
            r = _run(p, "", co=10, nghieng=True)
            for kieu, chu in (("begin", None), (None, f" {truong} "), ("end", None)):
                if kieu:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), kieu)
                else:
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = chu
                r._r.append(el)
            if truong == "PAGE":
                _run(p, "/", co=10, nghieng=True)


def dung_docx(khoi: list[Khoi], ra: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Cm, Pt

    doc = Document()
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.0)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(CO)
    doc.core_properties.title = "Bản kê khai công cụ AI, dữ liệu, API, thư viện — LiveLift"
    doc.core_properties.author = "Đội thi LiveLift"
    doc.core_properties.comments = "Dựng từ 05-BAN-KE-KHAI.md bằng ke_khai/dung_ke_khai.py"

    cho_ky = False
    for k in khoi:
        if k.loai == "chi_thi":
            if k.chu == "QUOC-HIEU":
                _quoc_hieu(doc)
            elif k.chu == "BANG-KY":
                cho_ky = True
                if doc.paragraphs:  # dòng "…, ngày … tháng …" đi cùng trang với bảng ký
                    doc.paragraphs[-1].paragraph_format.keep_with_next = True
            elif k.chu == "NGAT-TRANG":
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        if k.loai == "tieu_de":
            if k.cap == 1:
                _doan(doc, k.chu, co=15, dam=True, canh="giua", truoc=4, sau=10, giu_voi_sau=True)
            elif k.cap == 2:
                _doan(doc, k.chu, co=13.5, dam=True, canh="trai", truoc=10, sau=4, giu_voi_sau=True)
            else:
                _doan(
                    doc,
                    k.chu,
                    co=13,
                    dam=True,
                    nghieng=True,
                    canh="trai",
                    truoc=6,
                    sau=3,
                    giu_voi_sau=True,
                )
        elif k.loai == "doan":
            _doan(doc, k.chu)
        elif k.loai in ("ds_cham", "ds_so"):
            for n, (cap, chu) in enumerate(k.muc, 1):
                dau = ("•" if cap == 0 else "–") if k.loai == "ds_cham" else f"{n}."
                p = _doan(doc, "", thut_trai=0.6 + 0.6 * cap, thut_dau=-0.5, sau=2)
                _run(p, dau + "\t")
                _chu_dinh_dang(p, chu)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.6 + 0.6 * cap))
        elif k.loai == "bang":
            _bang(doc, k, ky_ten=cho_ky)
            cho_ky = False
        elif k.loai == "trich":
            _doan(doc, k.chu, co=12, nghieng=True, thut_trai=0.8, sau=4)
        elif k.loai == "ma":
            for i, dong in enumerate(k.muc):
                p = _doan(
                    doc, "", canh="trai", thut_trai=0.5, sau=0, giu_voi_sau=i < len(k.muc) - 1
                )
                _run(p, dong or " ", co=11, ma=True)
            _doan(doc, "", sau=2)
        elif k.loai == "ke":
            continue
    _so_trang(doc)
    doc.save(str(ra))


# ------------------------------------------------------------------ kiểm tra
def kiem_md(khoi: list[Khoi]) -> list[str]:
    """Kiểm nội dung nguồn: ký hiệu Markdown lọt ra, ô ⬜ chưa điền."""
    loi: list[str] = []
    chu: list[str] = []
    for k in khoi:
        if k.loai in ("tieu_de", "doan", "trich"):
            chu.append(chu_tron(k.chu))
        elif k.loai in ("ds_cham", "ds_so"):
            chu.extend(chu_tron(c) for _cap, c in k.muc)
        elif k.loai == "bang":
            chu.extend(chu_tron(o) for h in k.muc for o in h)
    for c in chu:
        if "`" in c or "**" in c or "](" in c:
            loi.append(f"ký hiệu Markdown lọt ra: {c[:80]!r}")
        if "⬜" in c:
            loi.append(f"còn ô ⬜ chưa điền: {c[:80]!r}")
        if "⟦" in c:
            loi.append(f"còn chỗ chờ điền số: {c[:80]!r}")
    return loi


HO_SO = DAY.parent / "noi-dung.md"
_MA_COMMIT_MUC_13 = re.compile(r"nhánh `main` tại commit `([0-9a-f]{7,40})`")
# Câu tả kho TRƯỚC khi hợp nhất nhánh hoàn thiện hồ sơ vào `main` (kiểm độc lập 25/09/2026).
_CAU_TRUOC_HOP_NHAT: tuple[tuple[str, str], ...] = (
    (r"`main` tại commit `390027b`", "trạng thái mã nguồn còn ghi main = 390027b"),
    (r"\b56/56\b", "còn ghi 56/56 commit trên main (I.1, VII.1)"),
    (r"0 trên `main`", "I.2 còn ghi Opus 5.5 có 0 commit trên main"),
    (r"Chờ trưởng nhóm hợp nhất", "II còn ghi chờ hợp nhất"),
    (r"Trên nhánh hoàn thiện hồ sơ", "II còn tả thay đổi là của nhánh chưa hợp nhất"),
    (r"nhánh hoàn thiện hồ sơ[^|\n]{0,40}chưa hợp nhất", "còn ghi nhánh hoàn thiện chưa hợp nhất"),
    (r"Commit trên `main` \| 56 commit", "VII.1 còn ghi 56 commit trên main"),
)


def lech_trang_thai_kho(ho_so: str, ke_khai: str) -> list[str]:
    """Hồ sơ mục 13 đã ghi mã commit `main` (tức đã hợp nhất và đẩy lên) thì bản kê khai không
    được còn tả kho như trước khi hợp nhất, và phải ghi đúng mã commit đó. Mục 13 còn ô ⬜ thì
    không kiểm (trả rỗng). Trả danh sách lỗi, rỗng là khớp."""
    m = _MA_COMMIT_MUC_13.search(ho_so)
    if not m:
        return []
    loi = [
        f"hồ sơ mục 13 trỏ main {m.group(1)} nhưng bản kê khai {ly_do} ({len(thay)} chỗ)"
        for mau, ly_do in _CAU_TRUOC_HOP_NHAT
        if (thay := re.findall(mau, ke_khai))
    ]
    if m.group(1)[:7] not in ke_khai:
        loi.append(f"bản kê khai không ghi mã commit main {m.group(1)[:7]} của hồ sơ mục 13")
    return loi


def kiem_docx(tep: Path) -> list[str]:
    from docx import Document
    from docx.oxml.ns import qn

    loi = []
    d = Document(str(tep))
    for t in d.element.body.iter(qn("w:t")):
        s = t.text or ""
        if "`" in s or "**" in s or "](" in s:
            loi.append(f"ký hiệu Markdown trong .docx: {s[:80]!r}")
        if "⬜" in s:
            loi.append(f"ô ⬜ trong .docx: {s[:80]!r}")
    return loi


def xuat_pdf(docx: Path, pdf: Path) -> tuple[int | None, int | None]:
    """Word đếm trang rồi xuất PDF. Trả (số trang, số từ)."""
    d_ps = str(docx).replace("'", "''")
    p_ps = str(pdf).replace("'", "''")
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        "$w.DisplayAlerts = 0; "
        f"$d = $w.Documents.Open('{d_ps}', $false, $true); $d.Fields.Update() | Out-Null; "
        "$d.Repaginate(); $t = $d.ComputeStatistics(2); $u = $d.ComputeStatistics(0); "
        f"$d.SaveAs([ref]'{p_ps}', [ref]17); $d.Close($false); $w.Quit(); "
        'Write-Output "$t $u"'
    )
    try:
        out = subprocess.run(  # noqa: S603 — lệnh dựng từ hằng số + đường dẫn do script sinh
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],  # noqa: S607
            capture_output=True,
            text=True,
            timeout=300,
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        print(f"  Không gọi được Word: {e}")
        return None, None
    so = out.stdout.split()
    if len(so) < 2 or not so[0].isdigit() or not pdf.exists():
        print("  Word không xuất được PDF.", (out.stderr.strip().splitlines() or [""])[0])
        return None, None
    return int(so[0]), int(so[1])


def main(argv: list[str] | None = None) -> int:
    for s in (sys.stdout, sys.stderr):
        with contextlib.suppress(AttributeError, ValueError):
            s.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--nguon", type=Path, default=NGUON)
    ap.add_argument("--ra", type=Path, default=RA_MAC_DINH, help="tệp .docx kết quả")
    ap.add_argument("--khong-pdf", action="store_true", help="không gọi Word, chỉ .docx")
    ap.add_argument("--chi-kiem", action="store_true", help="chỉ kiểm tệp .md, không dựng")
    a = ap.parse_args(argv)

    nguon = a.nguon.read_text(encoding="utf-8")
    khoi = phan_tich(nguon)
    loi = kiem_md(khoi)
    if HO_SO.exists():
        loi += lech_trang_thai_kho(HO_SO.read_text(encoding="utf-8"), nguon)
    if loi:
        print("CHƯA DỰNG — nguồn còn lỗi:")
        for x in loi:
            print("  -", x)
        return 1
    if a.chi_kiem:
        print(f"Nguồn sạch: {len(khoi)} khối.")
        return 0
    _nap_docx()
    tam = a.ra.with_name(a.ra.stem + ".tam.docx")
    dung_docx(khoi, tam)
    loi = kiem_docx(tam)
    if loi:
        tam.unlink(missing_ok=True)
        print("CHƯA DỰNG — .docx còn lỗi:")
        for x in loi:
            print("  -", x)
        return 1
    if a.khong_pdf:
        os.replace(tam, a.ra)
        print(f"Đã dựng {a.ra} (không xuất PDF).")
        return 0
    pdf_tam = tam.with_suffix(".pdf")
    trang, tu = xuat_pdf(tam, pdf_tam)
    if trang is None:
        tam.unlink(missing_ok=True)
        print("CHƯA DỰNG — không xuất được PDF; giữ nguyên bản cũ.")
        return 1
    os.replace(tam, a.ra)
    os.replace(pdf_tam, a.ra.with_suffix(".pdf"))
    print(f"Đã dựng {a.ra} và {a.ra.with_suffix('.pdf')} — {trang} trang, {tu} từ.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
